"""
Geração das planilhas ``FNDE_<Grupo de Insumo>.xlsx``.

Cada arquivo tem:
    * Aba "Fornecedores": fornecedores que cumprem todas as regras.
    * Aba "Sem preço (referência)": atacadistas da região com catálogo, mas
      preço oculto/sob consulta — não servem para a coleta, ficam só como
      referência.
    * Aba "Pendentes de validação": têm site com preço, mas o CNPJ/CNAE
      ainda não pôde ser confirmado — não usar até validar.
    * Aba "Resumo": contagem por UF x meta e regras aplicadas.

Para os grupos de serviço (PASSAGEM, SAÚDE OCUPACIONAL) a aba de
fornecedores sai vazia, com a observação definida em ``config.py``.
"""

from __future__ import annotations

import logging
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from fgv_precos.config import (
    COLUNAS_SAIDA,
    GRUPOS_SERVICO,
    META_POR_UF,
    NOME_ARQUIVO_POR_GRUPO,
    OBSERVACAO_GRUPO_SERVICO,
    UFS_PRIORITARIAS,
)

log = logging.getLogger(__name__)

SITUACAO_VALIDO = "VALIDO"
SITUACAO_REFERENCIA = "REFERENCIA_SEM_PRECO"
SITUACAO_PENDENTE = "PENDENTE_VALIDACAO"   # site com preço, mas CNAE/CNPJ não confirmado
COLUNAS_LINK = ("Site Oficial", "Link (clicável)", "Fonte")

_COR_CABECALHO = PatternFill("solid", start_color="1F4E78")
_FONTE_CABECALHO = Font(name="Arial", bold=True, color="FFFFFF", size=10)
_FONTE_NORMAL = Font(name="Arial", size=10)
_FONTE_LINK = Font(name="Arial", size=10, color="0563C1", underline="single")
_LARGURAS = {"Estado": 8, "Cidade": 18, "Nome da Empresa": 34, "CNPJ": 20,
             "CNAE atacadista": 26, "Site Oficial": 30, "Link (clicável)": 40,
             "Telefone": 18, "E-mail Comercial": 30, "Endereço": 45,
             "Observações": 55, "Fonte": 40}


def _escrever_tabela(ws: Worksheet, df: pd.DataFrame) -> None:
    """Escreve cabeçalho + linhas, com links clicáveis e formatação."""
    for c, nome in enumerate(COLUNAS_SAIDA, start=1):
        cel = ws.cell(row=1, column=c, value=nome)
        cel.fill, cel.font = _COR_CABECALHO, _FONTE_CABECALHO
        cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = _LARGURAS.get(nome, 20)

    for r, (_, linha) in enumerate(df.iterrows(), start=2):
        for c, nome in enumerate(COLUNAS_SAIDA, start=1):
            valor = str(linha.get(nome, "") or "").strip()
            cel = ws.cell(row=r, column=c, value=valor or None)
            cel.alignment = Alignment(vertical="top", wrap_text=True)
            # "Fonte" pode ter várias URLs separadas por " | "; o link vai na 1ª.
            primeira_url = valor.split(" | ")[0]
            if nome in COLUNAS_LINK and primeira_url.startswith("http"):
                cel.hyperlink = primeira_url
                cel.font = _FONTE_LINK
            else:
                cel.font = _FONTE_NORMAL
    ws.freeze_panes = "A2"
    if len(df):
        ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUNAS_SAIDA))}{len(df) + 1}"


def _escrever_resumo(ws: Worksheet, grupo: str, validos: pd.DataFrame,
                     observacao_extra: str = "") -> None:
    linhas: list[list[str]] = [
        ["Grupo de Insumo", grupo],
        ["Gerado em", date.today().strftime("%d/%m/%Y")],
        [],
        ["UF da cesta", "Sediados na UF", "De outras UFs", "Total",
         f"Meta (mínimo {META_POR_UF} por UF)"],
    ]
    atende = validos.get("Atende UF", pd.Series(dtype=str)).fillna("")
    estado = validos.get("Estado", pd.Series(dtype=str)).fillna("")
    # Grupos de serviço não têm meta de fornecedor atacadista.
    for uf in (() if grupo in GRUPOS_SERVICO else UFS_PRIORITARIAS):
        # "Atende UF" pode ser "GO", "DF" ou "GO/DF" (fornecedor de fora que completa ambas).
        atende_uf = atende.str.split("/").apply(lambda ufs, u=uf: u in {x.strip() for x in ufs})
        locais = int((atende_uf & (estado == uf)).sum())
        de_fora = int((atende_uf & (estado != uf)).sum())
        total = locais + de_fora
        linhas.append([uf, str(locais), str(de_fora), str(total),
                       "OK" if total >= META_POR_UF else f"faltam {META_POR_UF - total}"])
    linhas += [
        [],
        ["Regras aplicadas"],
        ["1. CNAE da Divisão 46 (atacado), principal ou secundário, conferido no CNPJ."],
        ["2. Site próprio com catálogo e preço visível sem login (marketplaces excluídos)."],
        ["3. Prioridade GO e DF; fornecedor de outra UF só completa a meta e "
         "traz a UF atendida em 'Observações'."],
        ["4. Nenhum dado inventado: campo não encontrado fica em branco."],
    ]
    if observacao_extra:
        linhas += [[], ["Observação"], [observacao_extra]]
    for linha in linhas:
        ws.append(linha)
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 22
    for col in ("C", "D", "E"):
        ws.column_dimensions[col].width = 26
    ws["A1"].font = Font(name="Arial", bold=True)


def exportar_grupo(grupo: str, df_grupo: pd.DataFrame, pasta_saida: Path) -> Path:
    """Gera o .xlsx de um Grupo de Insumo e devolve o caminho."""
    pasta_saida.mkdir(parents=True, exist_ok=True)
    nome_arquivo = NOME_ARQUIVO_POR_GRUPO.get(grupo, f"FNDE_{grupo.replace('/', '_')}.xlsx")
    caminho = pasta_saida / nome_arquivo

    validos = df_grupo[df_grupo["Situação"] == SITUACAO_VALIDO]
    referencia = df_grupo[df_grupo["Situação"] == SITUACAO_REFERENCIA]
    pendentes = df_grupo[df_grupo["Situação"] == SITUACAO_PENDENTE]

    wb = Workbook()
    ws_forn = wb.active
    ws_forn.title = "Fornecedores"
    _escrever_tabela(ws_forn, validos)
    if grupo in GRUPOS_SERVICO:
        ws_forn.cell(row=3, column=1, value=OBSERVACAO_GRUPO_SERVICO[grupo]).font = _FONTE_NORMAL

    if not referencia.empty:
        _escrever_tabela(wb.create_sheet("Sem preço (referência)"), referencia)
    if not pendentes.empty:
        _escrever_tabela(wb.create_sheet("Pendentes de validação"), pendentes)

    _escrever_resumo(wb.create_sheet("Resumo"), grupo, validos,
                     OBSERVACAO_GRUPO_SERVICO.get(grupo, ""))
    try:
        wb.save(caminho)
    except PermissionError as exc:
        # Erro comum no Windows: o arquivo está aberto no Excel.
        raise PermissionError(f"Feche o arquivo no Excel e rode de novo: {caminho}") from exc
    log.info("Planilha gerada: %s (%d válidos, %d referência, %d pendentes)",
             caminho.name, len(validos), len(referencia), len(pendentes))
    return caminho


def exportar_todos(df_curado: pd.DataFrame, pasta_saida: Path) -> list[Path]:
    """Gera uma planilha para cada Grupo de Insumo configurado."""
    gerados: list[Path] = []
    for grupo in NOME_ARQUIVO_POR_GRUPO:
        df_grupo = df_curado[df_curado["Grupo de Insumo"] == grupo]
        gerados.append(exportar_grupo(grupo, df_grupo, pasta_saida))
    return gerados
