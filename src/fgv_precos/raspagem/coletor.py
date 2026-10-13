"""
Orquestra a coleta: para cada insumo, percorre os fornecedores do grupo até
juntar 5 preços com frete para GO.

Regras (ver ``config_raspagem.py``):
    * Ordem dos fornecedores = ordem da aba "Fornecedores" do
      FNDE_<Grupo>.xlsx (os de GO/DF vêm primeiro).
    * Produto correspondente sem preço visível (login, sob consulta ou motivo
      não identificado) -> linha para coleta manual com o link na coluna P e
      o motivo em OBS FRETE; não conta entre os 5 preços e o próximo
      fornecedor é tentado. Só o produto claramente esgotado é ignorado.
    * Produto aceito: melhor pontuação de ``descricao.pontuar`` (> 0),
      disponível e com preço. Links da lista de produtos recusados na revisão
      (FNDE/Centro Oeste/produtos_recusados.csv) nunca são escolhidos de novo
      para o mesmo insumo.
    * Frete que exige login/cadastro -> linha na coluna P (link para coleta
      manual), com "Frete só com login" em OBS FRETE; não conta entre os 5.
    * Frete que não pode ser calculado por outro motivo -> fornecedor
      descartado.

Uso (o navegador fica aberto entre os insumos)::

    with Coletor(headless=True) as coletor:
        for insumo in insumos:
            coletas, resumo = coletor.coletar_insumo(insumo)
"""

from __future__ import annotations

import csv
import logging
import shutil
from dataclasses import replace
from datetime import date
from functools import cache
from pathlib import Path
from urllib.parse import urlparse

import openpyxl
from playwright.sync_api import Error as PlaywrightError
from playwright.sync_api import Page, sync_playwright

from fgv_precos.config import (
    ARQ_INSUMOS_PADRAO,
    DIR_LOGS,
    DIR_SAIDA_FORNECEDORES,
    NOME_ARQUIVO_POR_GRUPO,
)
from fgv_precos.fornecedores.cnae import somente_digitos
from fgv_precos.fornecedores.insumos import ler_insumos
from fgv_precos.precos import float_para_br
from fgv_precos.raspagem.config_raspagem import (
    ARGS_NAVEGADOR,
    ARQ_PRODUTOS_RECUSADOS,
    CANAIS_NAVEGADOR,
    CEPS,
    DIR_PRINTS,
    FRETE_ALERTA_MULTIPLO,
    PASTA_PRINT_POR_GRUPO,
    PRECOS_POR_UF,
    TIMEOUT_NAVEGACAO_MS,
    USER_AGENT_NAVEGADOR,
    VIEWPORT,
)
from fgv_precos.raspagem.descricao import Insumo, interpretar, pontuar
from fgv_precos.raspagem.lojas import loja_para
from fgv_precos.raspagem.modelos import Coleta, Fornecedor, Produto
from fgv_precos.raspagem.plataformas.base import MOTIVO_FRETE_LOGIN, LojaBase

log = logging.getLogger(__name__)

COLUNAS_RESUMO = ["ID FGV", "Grupo", "UF", "Preços coletados", "Faltam",
                  "Sem preço visível (coluna P)"]


@cache
def ler_fornecedores(grupo: str) -> tuple[Fornecedor, ...]:
    """Lê a aba 'Fornecedores' do FNDE_<Grupo>.xlsx (uma vez por execução)."""
    arquivo = DIR_SAIDA_FORNECEDORES / NOME_ARQUIVO_POR_GRUPO[grupo]
    if not arquivo.exists():
        log.error("Planilha de fornecedores não encontrada: %s", arquivo)
        return ()
    wb = openpyxl.load_workbook(arquivo, read_only=True, data_only=True)
    try:
        if "Fornecedores" not in wb.sheetnames:
            log.error("Aba 'Fornecedores' ausente em %s", arquivo.name)
            return ()
        linhas = list(wb["Fornecedores"].iter_rows(values_only=True))
    finally:
        wb.close()   # read_only mantém o arquivo aberto (e travado no Windows) até fechar
    if not linhas:
        return ()
    cab = [str(c or "").strip() for c in linhas[0]]
    if "Nome da Empresa" not in cab or "Site Oficial" not in cab:
        log.error("Colunas 'Nome da Empresa'/'Site Oficial' ausentes em %s", arquivo.name)
        return ()
    idx = {nome: cab.index(nome) for nome in ("Nome da Empresa", "CNPJ", "Site Oficial", "Estado")
           if nome in cab}

    def campo(linha: tuple, nome: str) -> str:
        return str(linha[idx[nome]] or "").strip() if nome in idx else ""

    saida = []
    for linha in linhas[1:]:
        nome, site = campo(linha, "Nome da Empresa"), campo(linha, "Site Oficial")
        if nome and not site:
            log.info("Fornecedor sem site ignorado: %s", nome)
        if nome and site:
            saida.append(Fornecedor(nome=nome, cnpj=campo(linha, "CNPJ"), site=site,
                                    estado=campo(linha, "Estado")))
    log.info("Grupo %s: %d fornecedores com site.", grupo, len(saida))
    return tuple(saida)


def fornecedor_combina(fornecedor: Fornecedor, filtro: str) -> bool:
    """True se o filtro estiver vazio, contido no nome ou igual ao CNPJ (só dígitos)."""
    if not filtro:
        return True
    digitos = somente_digitos(filtro)
    return (filtro.casefold() in fornecedor.nome.casefold()
            or (len(digitos) == 14 and digitos == somente_digitos(fornecedor.cnpj)))


def normalizar_url(url: str) -> str:
    """Endereço sem "www.", sem parâmetros e sem barra final, para comparar links."""
    partes = urlparse((url or "").strip())
    return f"{partes.netloc.lower().removeprefix('www.')}{partes.path.rstrip('/')}"


def ler_produtos_recusados(caminho: Path = ARQ_PRODUTOS_RECUSADOS) -> dict[str, set[str]]:
    """{ID FGV: links recusados na revisão}. Arquivo ausente = nenhum recusado."""
    if not caminho.exists():
        return {}
    recusados: dict[str, set[str]] = {}
    with open(caminho, encoding="utf-8-sig", newline="") as fh:
        for linha in csv.DictReader(fh, delimiter=";"):
            id_fgv, link = str(linha.get("ID FGV") or "").strip(), linha.get("Link") or ""
            if id_fgv and link.strip():
                recusados.setdefault(id_fgv, set()).add(normalizar_url(link))
    log.info("Produtos recusados na revisão: %d links em %d insumos.",
             sum(map(len, recusados.values())), len(recusados))
    return recusados


def ler_insumos_selecionados(ids: list[str] | None,
                             caminho: Path = ARQ_INSUMOS_PADRAO) -> list[Insumo]:
    df = ler_insumos(caminho)
    if ids:
        df = df[df["ID FGV"].astype(str).isin([str(i) for i in ids])]
        faltando = set(map(str, ids)) - set(df["ID FGV"].astype(str))
        if faltando:
            log.warning("IDs não encontrados na planilha de insumos: %s", sorted(faltando))
    return [interpretar(r["ID FGV"], r["Grupo de Insumo"], r["Descrição"], r.get("Embalagem", ""))
            for _, r in df.iterrows()]


def abrir_navegador(pw, headless: bool):
    """Abre o primeiro navegador disponível da lista ``CANAIS_NAVEGADOR``."""
    ultimo_erro = None
    for canal in CANAIS_NAVEGADOR:
        opcoes = {"headless": headless, "args": ARGS_NAVEGADOR}
        if canal:
            opcoes["channel"] = canal
        try:
            navegador = pw.chromium.launch(**opcoes)
            log.info("Navegador usado: %s", canal or "chromium do Playwright")
            return navegador
        except PlaywrightError as exc:
            ultimo_erro = exc
            log.info("Navegador '%s' indisponível: %s", canal, str(exc).splitlines()[0])
    raise RuntimeError(f"Nenhum navegador pôde ser aberto: {ultimo_erro}")


def _publicar_prints(temporaria: Path, pasta: Path, id_fgv: str,
                     coletas: list[Coleta]) -> None:
    """Troca os prints antigos do insumo pelos novos, só depois da coleta dele.

    Assim, se a execução for interrompida no meio, os prints antigos (que
    correspondem à planilha ainda não regravada) continuam no lugar.
    """
    for antigo in pasta.glob(f"{id_fgv}_*.png"):
        try:
            antigo.unlink()
        except OSError as exc:
            log.warning("Não foi possível apagar o print antigo %s: %s", antigo.name, exc)
    for coleta in coletas:
        if coleta.caminho_print:
            origem = Path(coleta.caminho_print)
            destino = pasta / origem.name
            origem.replace(destino)
            coleta.caminho_print = str(destino)
    shutil.rmtree(temporaria, ignore_errors=True)
    if temporaria.exists():
        log.debug("Pasta temporária não pôde ser apagada (OneDrive?): %s", temporaria)


class Coletor:
    """Mantém um navegador aberto e coleta insumo a insumo."""

    def __init__(self, headless: bool = True) -> None:
        self.headless = headless
        self.data_coleta = date.today().strftime("%d/%m/%Y")
        self._pw = self._navegador = self._contexto = None
        self.page: Page | None = None
        self.recusados = ler_produtos_recusados()

    # ------------------------------------------------------- ciclo de vida
    def __enter__(self) -> Coletor:
        self._pw = sync_playwright().start()
        self._navegador = abrir_navegador(self._pw, self.headless)
        self._contexto = self._navegador.new_context(
            viewport=VIEWPORT, locale="pt-BR", timezone_id="America/Sao_Paulo",
            user_agent=USER_AGENT_NAVEGADOR)
        self._contexto.set_default_timeout(TIMEOUT_NAVEGACAO_MS)
        self.page = self._contexto.new_page()
        return self

    def __exit__(self, *_exc) -> None:
        for recurso in (self._contexto, self._navegador):
            if recurso is not None:
                try:
                    recurso.close()
                except PlaywrightError as exc:
                    log.debug("Erro ao fechar o navegador: %s", exc)
        if self._pw is not None:
            self._pw.stop()

    # ------------------------------------------------------------- coleta
    def coletar_insumo(self, insumo: Insumo,
                       fornecedor_filtro: str = "") -> tuple[list[Coleta], list[dict]]:
        """Coleta um insumo e devolve (linhas da planilha, linhas do resumo).

        ``fornecedor_filtro`` (opcional) limita a coleta aos fornecedores cujo
        nome contém esse texto ou cujo CNPJ é esse (ex.: "singular").
        """
        if self.page is None:
            raise RuntimeError("Use o Coletor dentro de um bloco 'with'.")
        if insumo.grupo not in NOME_ARQUIVO_POR_GRUPO:
            log.warning("Grupo sem planilha de fornecedores: %s", insumo.grupo)
            return [], []
        log.info("=== Insumo %s (%s): busca %s | medidas %s",
                 insumo.id_fgv, insumo.grupo, insumo.nomes_busca,
                 sorted(insumo.medidas_obrigatorias) or "-")
        pasta = DIR_PRINTS / PASTA_PRINT_POR_GRUPO[insumo.grupo]
        temporaria = pasta / f"_coletando_{insumo.id_fgv}"
        shutil.rmtree(temporaria, ignore_errors=True)   # sobra de execução interrompida
        temporaria.mkdir(parents=True)

        contagem = {uf: 0 for uf in CEPS}
        coletas: list[Coleta] = []
        fornecedores = [f for f in ler_fornecedores(insumo.grupo)
                        if fornecedor_combina(f, fornecedor_filtro)]
        if fornecedor_filtro:
            log.info("Fornecedores com '%s': %s", fornecedor_filtro,
                     [f.nome for f in fornecedores] or "nenhum")
        for fornecedor in fornecedores:
            if all(n >= PRECOS_POR_UF for n in contagem.values()):
                break
            # "zeusdobrasil.com.br / epizeus.com.br": tenta cada site até achar o produto.
            for site in [s.strip() for s in fornecedor.site.split(" / ") if s.strip()]:
                try:
                    novas = self._tentar_fornecedor(insumo, replace(fornecedor, site=site),
                                                    contagem, temporaria)
                except PlaywrightError as exc:
                    log.error("[%s] Erro de navegador: %s", fornecedor.nome, exc)
                    novas = []
                except Exception:   # noqa: BLE001 - um site com problema não pode parar a coleta
                    log.exception("[%s] Erro inesperado", fornecedor.nome)
                    novas = []
                coletas.extend(novas)
                if any(not c.link_aproveitavel for c in novas):
                    break           # achou preço neste site; não precisa do próximo

        _publicar_prints(temporaria, pasta, insumo.id_fgv, coletas)
        resumo = []
        for uf, n in contagem.items():
            manual = sum(1 for c in coletas if c.uf == uf and c.link_aproveitavel)
            nivel = logging.INFO if n >= PRECOS_POR_UF else logging.WARNING
            log.log(nivel, "Insumo %s - %s: %d de %d preços (%d para coleta manual).",
                    insumo.id_fgv, uf, n, PRECOS_POR_UF, manual)
            resumo.append({"ID FGV": insumo.id_fgv, "Grupo": insumo.grupo, "UF": uf,
                           "Preços coletados": n, "Faltam": max(PRECOS_POR_UF - n, 0),
                           "Sem preço visível (coluna P)": manual})
        return coletas, resumo

    def _coleta(self, insumo: Insumo, uf: str, fornecedor: Fornecedor, produto: Produto,
                **extras) -> Coleta:
        return Coleta(insumo.id_fgv, insumo.grupo, insumo.descricao, uf, self.data_coleta,
                      fornecedor, produto, **extras)

    def _diagnosticar_frete(self, loja: LojaBase, insumo: Insumo) -> None:
        """Registra o que a página mostrou quando o frete não pôde ser lido:
        o texto do bloco do frete no log e um print em logs/diagnostico_frete/."""
        antes, depois = (" | ".join(" ".join(v.split()) for v in t.splitlines() if v.strip())
                         for t in loja.texto_frete)
        log.info("[%s] Bloco do frete ANTES: %s", loja.fornecedor.nome, antes[:1500] or "-")
        log.info("[%s] Bloco do frete DEPOIS: %s", loja.fornecedor.nome, depois[:1500] or "-")
        pasta = DIR_LOGS / "diagnostico_frete"
        pasta.mkdir(parents=True, exist_ok=True)
        caminho = pasta / f"{date.today():%Y%m%d}_{insumo.id_fgv}_{urlparse(loja.base).netloc}.png"
        if loja.tirar_print(str(caminho)):
            log.info("[%s] Print do diagnóstico: %s", loja.fornecedor.nome, caminho)

    def _escolher_produto(self, loja: LojaBase, insumo: Insumo) -> Produto | None:
        """Busca no site e devolve o primeiro produto aderente e disponível."""
        nome_loja = loja.fornecedor.nome
        candidatos = []
        recusados = getattr(self, "recusados", {}).get(insumo.id_fgv, set())
        # Tenta cada alternativa do nome (e, por fim, só a palavra principal).
        for termo in insumo.termos_de_busca:
            achados = loja.buscar(termo)
            for c in [c for c in achados if normalizar_url(c.url) in recusados]:
                log.info("[%s] Descartado '%s' (recusado na revisão)", nome_loja, c.nome)
            achados = [c for c in achados if normalizar_url(c.url) not in recusados]
            for c in achados:
                c.pontuacao = pontuar(insumo, c.nome)
            candidatos = sorted((c for c in achados if c.pontuacao > 0),
                                key=lambda c: c.pontuacao, reverse=True)
            log.info("[%s] Busca '%s': %d resultados, %d aderentes",
                     nome_loja, termo, len(achados), len(candidatos))
            if candidatos:
                break

        for cand in candidatos[:5]:
            produto = loja.ler_produto(cand.url)
            pontos = pontuar(insumo, produto.nome or cand.nome) if produto else 0
            if produto and normalizar_url(produto.url) in recusados:
                pontos = 0      # o link redirecionou para um produto recusado
            if produto and produto.disponivel and pontos > 0 and (
                    produto.preco_final or produto.motivo_sem_preco):
                log.info("[%s] Produto: %s | %s | aderência %.0f", nome_loja, produto.nome,
                         f"R$ {produto.preco_final:.2f}" if produto.preco_final
                         else produto.motivo_sem_preco, pontos)
                return produto
            log.info("[%s] Descartado '%s' (esgotado ou não aderente)", nome_loja, cand.nome)
        return None

    def _tentar_fornecedor(self, insumo: Insumo, fornecedor: Fornecedor,
                           contagem: dict[str, int], pasta: Path) -> list[Coleta]:
        loja = loja_para(self.page, fornecedor)
        loja.insumo = insumo
        log.info("[%s] %s (%s)", fornecedor.nome, loja.base, loja.plataforma)

        produto = self._escolher_produto(loja, insumo)
        if produto is None:
            return []
        if produto.motivo_sem_preco:
            # Sem preço visível: link para coleta manual (coluna P).
            log.info("[%s] %s: %s", fornecedor.nome, produto.motivo_sem_preco, produto.url)
            return [self._coleta(insumo, next(iter(CEPS)), fornecedor, produto, frete=None,
                                 link_aproveitavel=produto.url)]

        # Frete + print para cada UF que ainda precisa de preço.
        coletas = []
        for uf, cep in CEPS.items():
            if contagem[uf] >= PRECOS_POR_UF:
                continue
            if self.page.url != produto.url and not loja.abrir(produto.url):
                continue
            frete = loja.calcular_frete(cep, uf)
            if not frete.ok and frete.motivo == MOTIVO_FRETE_LOGIN:
                # Regra (09/10/2026): frete/carrinho que exige login vai para a
                # coluna P ("Link aproveitável"), para coleta manual; não conta
                # entre os 5 preços e o próximo fornecedor é tentado.
                motivo = ("Frete só com login (preço na página: "
                          f"R$ {float_para_br(produto.preco_final)})")
                log.info("[%s] %s: %s", fornecedor.nome, motivo, produto.url)
                coletas.append(self._coleta(insumo, uf, fornecedor,
                                            replace(produto, motivo_sem_preco=motivo),
                                            frete=None, link_aproveitavel=produto.url))
                continue
            if not frete.ok:
                log.info("[%s] Descartado para %s: %s", fornecedor.nome, uf, frete.motivo)
                self._diagnosticar_frete(loja, insumo)
                continue
            ordem = contagem[uf] + 1
            caminho = pasta / f"{insumo.id_fgv}_{uf}_{ordem}.png"
            if not loja.tirar_print(str(caminho)):
                continue
            contagem[uf] = ordem
            coletas.append(self._coleta(insumo, uf, fornecedor, produto, frete=frete,
                                        caminho_print=str(caminho)))
            log.info("[%s] %s OK: frete %s -> %s", fornecedor.nome, uf,
                     "grátis" if frete.gratis else f"R$ {frete.valor:.2f} ({frete.modalidade})",
                     caminho.name)
            if frete.valor and frete.valor > FRETE_ALERTA_MULTIPLO * produto.preco_final:
                log.warning("[%s] CONFERIR: frete R$ %.2f é mais de %.0fx o preço R$ %.2f (%s)",
                            fornecedor.nome, frete.valor, FRETE_ALERTA_MULTIPLO,
                            produto.preco_final, caminho.name)
        return coletas
