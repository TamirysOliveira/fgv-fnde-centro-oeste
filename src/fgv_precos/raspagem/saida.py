"""
Gravação do resultado em Excel (``Entrega/FNDE_output.xlsx``), no layout de
FNDE/Centro Oeste/FNDE_output.csv (o cabeçalho é lido do próprio modelo).

Cada gravação sobrescreve as linhas dos insumos coletados e mantém as linhas
dos demais insumos já presentes no arquivo. A gravação é feita ao fim de cada
insumo, para que uma interrupção não perca o que já foi coletado. Se a saída
anterior não puder ser lida, ela é renomeada (``*_ilegivel_<data>``) em vez
de ser sobrescrita.

Mapeamento (confirmado com a equipe):
    A Categoria          = Grupo de Insumo
    B Código FGV         = ID FGV
    C Item               = Descrição completa do insumo
    D UF                 = UF de onde a mercadoria sai: coluna "Estado" da planilha
                           de fornecedores (UF do estabelecimento do CNPJ da
                           coluna G). GO é só o destino do frete (CEP 74933-130)
                           e não vai para esta coluna. Regra de 13/10/2026.
    E Data coleta        = data da execução
    F Nome Empresa       = da planilha de fornecedores
    G CNPJ               = da planilha de fornecedores
    H FONTE              = link (clicável) da página do produto de onde saiu o print
    I PRODUTO PESQUISADO = nome do produto no site
    J Preço PRODUTO      = preço visível na página (sem cupom/PIX)
    K VALOR DESCONTO     = vazio (descontos não são considerados)
    L OBS Desconto       = vazio
    M PREÇO FINAL        = igual ao Preço PRODUTO
    N VALOR FRETE        = opção de entrega mais barata; vazio se grátis
    O OBS FRETE          = "Frete grátis..." ou a modalidade escolhida + CEP;
                           nas linhas da coluna P, o motivo da falta de preço
    P Link aproveitável  = só quando o produto correspondente não mostra preço
                           (login, sob consulta ou motivo não identificado):
                           link do produto para coleta manual; colunas de
                           preço/frete ficam vazias
"""

from __future__ import annotations

import csv
import logging
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from fgv_precos.raspagem.config_raspagem import (
    ARQ_MODELO_SAIDA,
    ARQ_RESUMO,
    ARQ_SAIDA,
    CEPS,
    CODIFICACAO_SAIDA,
    SEPARADOR_SAIDA,
)
from fgv_precos.raspagem.modelos import Coleta
from fgv_precos.raspagem.uf_origem import normalizar_uf

log = logging.getLogger(__name__)

N_COLUNAS_ESPERADAS = 16
COL_CODIGO, COL_FONTE, COL_LINK_APROVEITAVEL = 2, 8, 16    # B, H e P (1 = A)
COLUNAS_VALOR = (10, 11, 13, 14)                            # J, K, M, N: números no Excel
COLUNAS_QUEBRA_TEXTO = (3, 9, 15)                           # C, I, O
LARGURAS = [18, 11, 45, 5, 11, 32, 20, 40, 40, 13, 12, 12, 12, 12, 45, 40]
_FONTE_LINK = Font(name="Arial", size=10, color="0563C1", underline="single")
_FONTE = Font(name="Arial", size=10)


def colunas_modelo(caminho: Path = ARQ_MODELO_SAIDA) -> list[str]:
    """Lê o cabeçalho do modelo; assim a saída segue sempre o arquivo oficial.

    Raises:
        ValueError: se o modelo não tiver as 16 colunas (A a P) que o
            mapeamento acima pressupõe; gravar assim desalinharia as colunas.
    """
    with open(caminho, encoding=CODIFICACAO_SAIDA, newline="") as fh:
        cab = next(csv.reader(fh, delimiter=SEPARADOR_SAIDA))
    if len(cab) != N_COLUNAS_ESPERADAS:
        raise ValueError(f"{caminho.name} tem {len(cab)} colunas (esperado "
                         f"{N_COLUNAS_ESPERADAS}): {cab}. Ajuste o mapeamento em saida.py.")
    return cab


def _obs_frete(c: Coleta) -> str:
    f = c.frete
    cep = f"{CEPS[c.uf][:5]}-{CEPS[c.uf][5:]}"
    if f.gratis and f.modalidade.startswith("Frete grátis informado"):
        return f"{f.modalidade} (sem calculadora de CEP no site)"
    if f.gratis:
        return f"Frete grátis para o CEP {cep}"
    return f"Opção mais barata para o CEP {cep}: {f.modalidade}"


def uf_origem(c: Coleta) -> str:
    """Coluna D: UF do fornecedor (de onde a mercadoria sai), não a UF de destino.

    ``c.uf`` é a UF da cesta (destino do frete, GO) e continua sendo usada para o
    CEP, os prints e o resumo; ela não vai para a coluna D.
    """
    uf = normalizar_uf(c.fornecedor.estado)
    if not uf:
        log.warning("CONFERIR: %s sem UF válida na coluna 'Estado' da planilha de "
                    "fornecedores (%r); a coluna D (UF) fica vazia.",
                    c.fornecedor.nome, c.fornecedor.estado)
    return uf


def _linha(c: Coleta) -> list:
    base = [c.grupo, c.insumo_id, c.descricao, uf_origem(c), c.data, c.fornecedor.nome,
            c.fornecedor.cnpj]
    p = c.produto
    if c.link_aproveitavel:
        # Sem preço visível: nada de preço/frete; motivo em OBS FRETE, link na coluna P.
        return base + ["", p.nome, None, None, None, None, None, p.motivo_sem_preco,
                       c.link_aproveitavel]
    if c.frete is None or p.preco_final is None:
        raise ValueError(f"Coleta do insumo {c.insumo_id} sem preço/frete e sem link (coluna P).")
    preco = round(p.preco_final, 2)
    valor_frete = None if c.frete.gratis else round(c.frete.valor, 2)
    return base + [p.url, p.nome, preco, None, None, preco, valor_frete, _obs_frete(c), None]


def preservar_ilegivel(arquivo: Path, exc: Exception) -> None:
    """Renomeia um arquivo de saída que não pôde ser lido, para não perdê-lo."""
    carimbo = datetime.now().strftime("%Y%m%d_%H%M%S")
    copia = arquivo.with_name(f"{arquivo.stem}_ilegivel_{carimbo}{arquivo.suffix}")
    arquivo.replace(copia)
    log.error("Saída anterior ilegível (%s). Ela foi preservada como %s; os insumos "
              "dela precisam ser conferidos ou coletados de novo.", exc, copia.name)


def _ler_linhas_existentes(arquivo: Path, n_colunas: int) -> list[list]:
    """Linhas (sem cabeçalho) de uma saída anterior, para preservar outros insumos."""
    if not arquivo.exists():
        return []
    try:
        wb = load_workbook(arquivo, read_only=True)
        try:
            return [list(r[:n_colunas]) for r in wb.active.iter_rows(min_row=2, values_only=True)
                    if any(v not in (None, "") for v in r)]
        finally:
            wb.close()
    except PermissionError:
        raise PermissionError(f"Feche o arquivo no Excel e rode de novo: {arquivo}") from None
    except Exception as exc:  # noqa: BLE001 - qualquer falha de leitura: preserva o arquivo
        preservar_ilegivel(arquivo, exc)
        return []


def _escrever_linha(ws, valores: list, n_colunas: int) -> None:
    ws.append(valores)
    r = ws.max_row
    for col in range(1, n_colunas + 1):
        cel = ws.cell(row=r, column=col)
        cel.font = _FONTE
        cel.alignment = Alignment(vertical="top", wrap_text=col in COLUNAS_QUEBRA_TEXTO)
        if col in COLUNAS_VALOR and isinstance(cel.value, (int, float)):
            cel.number_format = "#,##0.00"
    for col in (COL_FONTE, COL_LINK_APROVEITAVEL):
        cel = ws.cell(row=r, column=col)
        if str(cel.value or "").startswith("http"):
            cel.hyperlink = cel.value
            cel.font = _FONTE_LINK


def escrever_planilha(linhas: list[list], cab: list[str], arquivo: Path) -> Path:
    """Escreve a planilha no layout do FNDE_output (cabeçalho, links, formatos)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "FNDE_output"
    ws.append(cab)
    for cel in ws[1]:
        cel.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cel.fill = PatternFill("solid", start_color="1F4E78")
        cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for linha in linhas:
        _escrever_linha(ws, linha, len(cab))
    for i, largura in enumerate(LARGURAS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = largura
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cab))}{ws.max_row}"
    try:
        wb.save(arquivo)
    except PermissionError as exc:
        raise PermissionError(f"Feche o arquivo no Excel e rode de novo: {arquivo}") from exc
    return arquivo


def ordenar_linhas(linhas: list[list]) -> list[list]:
    """Por código; dentro do insumo, preços completos antes das linhas da coluna P."""
    return sorted(linhas, key=lambda linha: (
        str(linha[COL_CODIGO - 1]), bool(linha[COL_LINK_APROVEITAVEL - 1])))


def gravar(coletas: list[Coleta], ids_coletados: list[str], arquivo: Path = ARQ_SAIDA,
           modelo: Path = ARQ_MODELO_SAIDA) -> Path:
    """Grava/atualiza ``Entrega/FNDE_output.xlsx``.

    As linhas dos insumos coletados SUBSTITUEM as linhas desses mesmos insumos
    na saída anterior; as linhas dos demais insumos são mantidas.
    """
    cab = colunas_modelo(modelo)
    arquivo.parent.mkdir(parents=True, exist_ok=True)

    ids = {str(i) for i in ids_coletados}
    mantidas = [linha for linha in _ler_linhas_existentes(arquivo, len(cab))
                if str(linha[COL_CODIGO - 1]) not in ids]
    novas = [_linha(c) for c in coletas]
    todas = ordenar_linhas(mantidas + novas)

    escrever_planilha(todas, cab, arquivo)
    log.info("Saída gravada: %s (%d linhas novas, %d de outros insumos mantidas)",
             arquivo.name, len(novas), len(mantidas))
    return arquivo


def gravar_resumo(novo: pd.DataFrame, ids: list[str], arquivo: Path = ARQ_RESUMO) -> Path:
    """Atualiza o resumo da coleta: substitui os IDs coletados, mantém os demais."""
    arquivo.parent.mkdir(parents=True, exist_ok=True)
    if arquivo.exists():
        try:
            antigo = pd.read_csv(arquivo, sep=";", dtype=str, encoding="utf-8-sig")
            antigo = antigo[~antigo["ID FGV"].astype(str).isin([str(i) for i in ids])]
            novo = pd.concat([antigo, novo.astype(str)], ignore_index=True)
        except PermissionError:
            raise PermissionError(f"Feche o arquivo no Excel e rode de novo: {arquivo}") from None
        except (OSError, ValueError, KeyError) as exc:
            preservar_ilegivel(arquivo, exc)
    novo = novo.sort_values(["ID FGV", "UF"])
    try:
        novo.to_csv(arquivo, sep=";", index=False, encoding="utf-8-sig")
    except PermissionError as exc:
        raise PermissionError(f"Feche o arquivo no Excel e rode de novo: {arquivo}") from exc
    log.info("Resumo gravado: %s", arquivo.name)
    return arquivo
