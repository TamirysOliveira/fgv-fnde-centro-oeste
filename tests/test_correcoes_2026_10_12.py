"""Testes das correções de 12/10/2026 (revisão do código e da entrega de 10/10).

Cada teste reproduz um caso real encontrado na planilha entregue ou no log.
"""

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from fgv_precos.fornecedores import exportador
from fgv_precos.fornecedores.cnpj_api import montar_endereco
from fgv_precos.fornecedores.receita_dump import montar_campos_contato
from fgv_precos.fornecedores.site_checker import palavras_encontradas
from fgv_precos.precos import valores_de_preco
from fgv_precos.raspagem import coletor, saida
from fgv_precos.raspagem.descricao import interpretar, pontuar
from fgv_precos.raspagem.modelos import Coleta, Fornecedor, Frete, Produto, limpar_nome
from fgv_precos.raspagem.plataformas.base import LojaBase
from fgv_precos.raspagem.plataformas.woocommerce import conferir_carrinho

# ---------------------------------------------------------------- frete


def test_carrinho_woocommerce_so_com_o_produto():
    """Casa da Limpeza, 10/10: 8 itens (R$ 1.136,80) liberaram frete grátis."""
    assert conferir_carrinho(1, 5.90, 5.90) == ""
    assert "8 itens" in conferir_carrinho(8, 1136.80, 5.90)
    assert "diferente do preço" in conferir_carrinho(1, 1136.80, 5.90)


def test_retirada_com_rotulo_na_linha_anterior_nao_e_frete():
    """Singular, 10/10: "Retirada em loja" / "R$ 0,00" virou VALOR FRETE = 0."""
    texto = ("Retirada em loja\nR$ 0,00\n4 dia(s) útil(eis)\n"
             "Pac\nR$ 27,30\n6 dia(s) útil(eis)")
    f = LojaBase.interpretar_frete(texto)
    assert f.ok and not f.gratis and f.valor == 27.30
    assert "Retirada" not in f.modalidade


def test_retirada_com_rotulo_na_linha_seguinte_nao_e_frete():
    """Loja Integrada (BHZ): "R$ 0,00 1 dia útil" / "Retirar pessoalmente"."""
    texto = ("R$ 0,00 1 dia útil\nRetirar pessoalmente\n"
             "R$ 20,62 7 dias úteis\nPAC\nR$ 26,55 5 dias úteis\nJadlog Package")
    f = LojaBase.interpretar_frete(texto)
    assert f.valor == 20.62 and f.modalidade == "R$ 20,62 7 dias úteis PAC"


def test_zero_sem_retirada_e_frete_gratis():
    f = LojaBase.interpretar_frete("Entrega econômica\nR$ 0,00\n5 dias úteis")
    assert f.ok and f.gratis


def test_obs_frete_sem_botoes_da_pagina():
    """10/10: OBS FRETE trazia "Lista de Desejos", "Compartilhar"..."""
    antes = "Calcule o frete\nCalcular"
    depois = antes + "\nLista de Desejos\nR$ 31,54\n9 dias úteis\nCorreios PAC\nCompartilhar"
    f = LojaBase.interpretar_frete(depois, antes)
    assert f.valor == 31.54
    assert f.modalidade == "R$ 31,54 9 dias úteis Correios PAC"


# ----------------------------------------------------------- aderência


def test_so_a_palavra_principal_nao_basta():
    carro = interpretar("1063", "LIMPEZA PREDIAL", "CARRO FUNCIONAL PARA LIMPEZA - COMPLETO")
    assert pontuar(carro, "Carro Guia Macarico Record") == 0
    assert pontuar(carro, "Carro Funcional Bralimpia América") > 0
    placa = interpretar("1092", "X", "PLACA SINALIZADORA – PISO MOLHADO - DOBRÁVEL")
    assert pontuar(placa, "Limpador de Placa Solar 5L Deo") == 0
    solar = interpretar("1119", "X", "BLOQUEADOR/PROTETOR SOLAR FPS 30 OU 50 - EXTERNO")
    assert pontuar(solar, "Protetor Auricular Plug k10 c/ Cordão 13DB") == 0
    assert pontuar(solar, "Protetor Solar FPS 30 Sundown 120ml") > 0


def test_genero_plural_e_sinonimos():
    escova = interpretar("1076", "X", "ESCOVA PARA VASO SANITÁRIO - COM SUPORTE.")
    assert pontuar(escova, "Escova Sanitária c/ Suporte") > 0
    disp = interpretar("1072", "X", "DISPENSADOR PARA SABONETE LÍQUIDO/ÁLCOOL GEL - DOSADOR")
    assert pontuar(disp, "Dispenser para Sabonete Líquido 800ml") > 0
    assert pontuar(disp, "Álcool Gel 70° 5L - Barbarex") == 0
    flanela = interpretar("1022", "X", "FLANELA PARA LIMPEZA - ALGODÃO")
    assert pontuar(flanela, "Flanela Branca G 38x58cm") > 0     # "limpeza" é genérico


def test_barra_nao_gera_busca_por_palavra_solta():
    assert interpretar("1108", "X", "BOTA DE PVC/BORRACHA - CANO MÉDIO").nomes_busca == [
        "Bota De Pvc", "Bota De Borracha"]


# --------------------------------------------------- nomes e preços


def test_nome_com_entidade_html():
    assert limpar_nome("Camiseta com Prote&#231;&#227;o UV ") == "Camiseta com Proteção UV"


def test_preco_ignora_parcela_e_pix():
    assert valores_de_preco("R$ 49,90 ou 10x de R$ 4,99 sem juros") == [49.90]
    assert valores_de_preco("R$ 49,90 R$ 47,40 no PIX") == [49.90]
    assert valores_de_preco("De R$ 59,90 por R$ 49,90") == [59.90, 49.90]


# ------------------------------------------------------ fornecedores


def test_palavra_chave_inteira():
    achadas = palavras_encontradas("Av. Rodovia GO-060, botão de contato; Rodo 40 cm e balde")
    assert achadas["LIMPEZA PREDIAL"] == {"rodo", "balde"}
    assert "bota" not in achadas["EPI / UNIFORME"]


def test_campos_contato_com_e_sem_linhas():
    colunas = ["cnpj_basico", "cnpj_ordem", "cnpj_dv", "ddd1", "telefone1", "tipo_logradouro",
               "logradouro", "numero", "complemento", "bairro", "cep", "email"]
    vazio = montar_campos_contato(pd.DataFrame(columns=colunas, dtype=str))
    assert vazio.empty and "telefone" in vazio.columns
    df = pd.DataFrame([["07964337", "0001", "68", "62", "39422020", "AV", "T-63", "2761", "",
                        "JD AMERICA", "", "Vendas@X.com "]], columns=colunas)
    linha = montar_campos_contato(df).iloc[0]
    assert linha["cnpj"] == "07.964.337/0001-68"
    assert linha["telefone"] == "(62) 39422020"
    assert linha["endereco"] == "AV T-63 2761 JD AMERICA"        # sem " - CEP " vazio
    assert linha["email"] == "vendas@x.com"


def test_montar_endereco():
    dados = {"logradouro": "Rua A", "numero": 10, "complemento": None, "bairro": "Centro"}
    assert montar_endereco(dados, ("logradouro", "numero", "complemento", "bairro"), "74000000") \
        == "Rua A 10 Centro - CEP 74000000"


def test_exportar_grupo(tmp_path):
    df = pd.DataFrame([{"Grupo de Insumo": "LIMPEZA PREDIAL", "Situação": "VALIDO",
                        "Atende UF": "GO / DF", "Estado": "GO", "Nome da Empresa": "Loja X",
                        "Site Oficial": "https://x.com.br"}])
    arquivo = exportador.exportar_grupo("LIMPEZA PREDIAL", df, tmp_path)
    resumo = load_workbook(arquivo)["Resumo"]
    valores = [[c.value for c in linha] for linha in resumo.iter_rows()]
    assert ["GO", "1", "0", "1", "faltam 4"] in valores
    assert ["DF", "0", "1", "1", "faltam 4"] in valores        # "GO / DF" com espaços


# ------------------------------------------------------------ gravação


def _coleta(id_fgv: str, preco: float = 10.0) -> Coleta:
    fornecedor = Fornecedor("Loja", "00.000.000/0001-00", "https://loja.com.br")
    produto = Produto(f"Produto {id_fgv}", f"https://loja.com.br/{id_fgv}", preco)
    return Coleta(id_fgv, "LIMPEZA PREDIAL", "Desc", "GO", "12/10/2026", fornecedor, produto,
                  Frete(ok=True, valor=20.0, modalidade="R$ 20,00 PAC"))


@pytest.fixture
def modelo(tmp_path):
    caminho = tmp_path / "FNDE_output.csv"
    caminho.write_text(";".join(f"Col{i}" for i in range(16)) + "\n", encoding="cp1252")
    return caminho


def test_gravar_substitui_so_o_insumo_coletado(tmp_path, modelo):
    arquivo = tmp_path / "FNDE_output.xlsx"
    saida.gravar([_coleta("1000"), _coleta("1001")], ["1000", "1001"], arquivo, modelo)
    saida.gravar([_coleta("1000", preco=99.0)], ["1000"], arquivo, modelo)
    linhas = list(load_workbook(arquivo).active.iter_rows(min_row=2, values_only=True))
    assert [(r[1], r[9]) for r in linhas] == [("1000", 99.0), ("1001", 10.0)]


def test_saida_ilegivel_e_preservada(tmp_path, modelo):
    arquivo = tmp_path / "FNDE_output.xlsx"
    arquivo.write_bytes(b"isto nao e um xlsx")
    saida.gravar([_coleta("1000")], ["1000"], arquivo, modelo)
    assert len(list(tmp_path.glob("FNDE_output_ilegivel_*.xlsx"))) == 1
    assert arquivo.exists()


def test_modelo_com_colunas_erradas_interrompe(tmp_path):
    modelo = tmp_path / "FNDE_output.csv"
    modelo.write_text("A;B;C\n", encoding="cp1252")
    with pytest.raises(ValueError, match="3 colunas"):
        saida.colunas_modelo(modelo)


def test_resumo_mantem_outros_insumos(tmp_path):
    arquivo = tmp_path / "resumo.csv"
    linha = {"ID FGV": "1000", "Grupo": "G", "UF": "GO", "Preços coletados": 5, "Faltam": 0,
             "Sem preço visível (coluna P)": 0}
    saida.gravar_resumo(pd.DataFrame([linha, {**linha, "ID FGV": "1001"}]), ["1000", "1001"],
                        arquivo)
    saida.gravar_resumo(pd.DataFrame([{**linha, "Preços coletados": 3}]), ["1000"], arquivo)
    df = pd.read_csv(arquivo, sep=";", dtype=str, encoding="utf-8-sig")
    assert df.set_index("ID FGV")["Preços coletados"].to_dict() == {"1000": "3", "1001": "5"}


def test_prints_antigos_so_saem_depois_da_coleta(tmp_path):
    pasta, temporaria = tmp_path, tmp_path / "_coletando_1000"
    temporaria.mkdir()
    (pasta / "1000_GO_1.png").write_bytes(b"antigo1")
    (pasta / "1000_GO_2.png").write_bytes(b"antigo2")
    (pasta / "10000_GO_1.png").write_bytes(b"outro insumo")
    novo = temporaria / "1000_GO_1.png"
    novo.write_bytes(b"novo")
    c = _coleta("1000")
    c.caminho_print = str(novo)
    coletor._publicar_prints(temporaria, pasta, "1000", [c])
    assert sorted(p.name for p in pasta.glob("*.png")) == ["10000_GO_1.png", "1000_GO_1.png"]
    assert (pasta / "1000_GO_1.png").read_bytes() == b"novo"
    assert c.caminho_print == str(pasta / "1000_GO_1.png")
    assert not temporaria.exists()


def test_ler_fornecedores_uma_vez_e_fecha(tmp_path, monkeypatch):
    wb = Workbook()
    ws = wb.active
    ws.title = "Fornecedores"
    ws.append(["Estado", "Nome da Empresa", "CNPJ", "Site Oficial"])
    ws.append(["GO", "Loja A", "1", "https://a.com.br"])
    ws.append(["GO", "Loja sem site", "2", None])
    wb.save(tmp_path / "FNDE_LIMPEZA PREDIAL.xlsx")
    monkeypatch.setattr(coletor, "DIR_SAIDA_FORNECEDORES", tmp_path)
    coletor.ler_fornecedores.cache_clear()
    primeira = coletor.ler_fornecedores("LIMPEZA PREDIAL")
    assert [f.nome for f in primeira] == ["Loja A"]
    assert coletor.ler_fornecedores("LIMPEZA PREDIAL") is primeira     # lida uma vez só
    (tmp_path / "FNDE_LIMPEZA PREDIAL.xlsx").rename(tmp_path / "x.xlsx")   # arquivo foi fechado
    coletor.ler_fornecedores.cache_clear()


class _CampoComMascara:
    """Imita o campo da Loja do EPI: ao digitar, o 1º dígito vai para o fim."""

    def __init__(self, embaralha: bool = True, aceita_colar: bool = True):
        self.valor, self.embaralha, self.aceita_colar = "", embaralha, aceita_colar

    def click(self):
        pass

    def fill(self, texto):
        self.valor = texto if (self.aceita_colar or not texto) else ""

    def press_sequentially(self, texto, delay=0):
        self.valor = texto[1:] + texto[0] if self.embaralha else texto

    def input_value(self):
        return f"{self.valor[:5]}-{self.valor[5:]}" if len(self.valor) > 5 else self.valor


def test_cep_conferido_no_campo():
    """Singular, 10/10: 74933-130 virou 49331-307 e o frete saiu para Sergipe."""
    loja = LojaBase(page=None, fornecedor=Fornecedor("Loja", "", "https://loja.com.br"))
    assert loja.preencher_cep(_CampoComMascara(embaralha=False), "74933130") == ""
    campo = _CampoComMascara()                       # digitar falha, colar funciona
    assert loja.preencher_cep(campo, "74933130") == "" and campo.valor == "74933130"
    problema = loja.preencher_cep(_CampoComMascara(aceita_colar=False), "74933130")
    assert "em vez de 74933130" in problema          # nada funcionou: frete não vale


def test_frete_gratis_com_rotulo_ao_lado_da_retirada():
    """Fermarcs, 10/10: "Frete Grátis Goiânia" é entrega de verdade (1 a 2 dias)."""
    texto = ("Frete\tValor\tPrazo\n"
             "Retirada na Loja RETIRAR NA LOJA AV. CONTORNO\tFRETE GRÁTIS\t1 dia(as). Avisos:\n"
             "Braspress TRANPORTADORA BRASSPRESS\tR$ 135,29\t1 dia(as). Avisos:\n"
             "Retirar na loja\tA RETIRAR\tO pedido ficará disponível para retirada em até 2 horas\n"
             "Frete Grátis Goiânia\tFRETE GRÁTIS\t1 a 2 (14/10/2026 até 15/10/2026). Aviso:\n"
             "Transportadora Do Valle\tR$ 51,59\t2 a 4 dia(as). Avisos:")
    assert LojaBase.interpretar_frete(texto).gratis
    so_retirada = "\n".join(v for v in texto.splitlines() if "Goiânia" not in v)
    f = LojaBase.interpretar_frete(so_retirada)
    assert not f.gratis and f.valor == 51.59


def test_loja_que_nao_entrega_no_cep():
    """Maroma, 10/10: o "frete" de R$ 5,99 era o preço de um produto relacionado."""
    texto = ("Não foram encontradas formas de envio para o CEP informado.\n"
             "Produtos relacionados\nEsponja Mágica Bompack UN\nR$ 5,99\nVer mais")
    f = LojaBase.interpretar_frete(texto)
    assert not f.ok and "não entrega" in f.motivo


def test_filtro_de_fornecedor():
    singular = Fornecedor("SINGULAR COMERCIO DE EPI LTDA (Loja do EPI)", "20.838.307/0001-73",
                          "https://www.lojadoepi.com.br")
    assert coletor.fornecedor_combina(singular, "")
    assert coletor.fornecedor_combina(singular, "singular")
    assert coletor.fornecedor_combina(singular, "20.838.307/0001-73")
    assert not coletor.fornecedor_combina(singular, "bhz")


def test_obs_frete_nao_pega_o_rotulo_da_proxima_opcao():
    """Singular, 12/10: OBS FRETE terminava com "Sedex" (rótulo da opção seguinte)."""
    texto = ("Retirada em loja\nR$ 0,00\n4 dia(s) útil(eis)\nPac\nR$ 25,47\n7 dia(s) útil(eis)\n"
             "Sedex\nR$ 31,27\n3 dia(s) útil(eis)")
    f = LojaBase.interpretar_frete(texto)
    assert f.valor == 25.47 and f.modalidade == "Pac R$ 25,47 7 dia(s) útil(eis)"


# ------------------------------------------------- UF de origem (13/10/2026)


def test_coluna_d_traz_a_uf_do_fornecedor_e_nao_a_do_destino():
    """Entrega de 13/10: a coluna D vinha GO (destino do frete) em todas as linhas."""
    c = _coleta("1000")
    c.fornecedor.estado = "sp "
    linha = saida._linha(c)
    assert linha[3] == "SP"
    assert "74933-130" in linha[14]          # o frete continua calculado para GO
    c.fornecedor.estado = ""
    assert saida._linha(c)[3] == ""          # sem UF na planilha: vazio, nunca GO


def test_mapa_uf_por_cnpj_e_por_site(tmp_path):
    from fgv_precos.raspagem.uf_origem import ler_mapa_uf

    wb = Workbook()
    ws = wb.active
    ws.title = "Fornecedores"
    ws.append(["Estado", "Cidade", "Nome da Empresa", "CNPJ", "CNAE atacadista", "Site Oficial",
               "Link (clicável)"])
    ws.append(["SP", "Barueri", "Gimba", "54.651.716/0011-50", "", "https://www.gimba.com.br", ""])
    ws.append(["SC", "Blumenau", "Zeus", "82.699.588/0001-88", "", "", "https://epizeus.com.br"])
    ws.append(["RS", "Porto Alegre", "Repetida", "11.111.111/0001-11", "", "https://a.com.br", ""])
    ws.append(["PR", "Curitiba", "Repetida", "11.111.111/0001-11", "", "https://a.com.br", ""])
    wb.create_sheet("Resumo").append(["Grupo de Insumo", "X"])
    arquivo = tmp_path / "FNDE_X.xlsx"
    wb.save(arquivo)

    mapa = ler_mapa_uf([arquivo, tmp_path / "nao_existe.xlsx"])
    assert mapa.uf("54651716001150") == "SP"                                   # pelo CNPJ
    assert mapa.uf("", "https://www.epizeus.com.br/luva?x=1") == "SC"          # pelo site
    assert mapa.uf(None, "", "https://loja.gimba.com.br/p") == "SP"            # subdomínio
    assert mapa.uf("11.111.111/0001-11") == ""                                 # UFs diferentes
    assert mapa.uf("99.999.999/0001-99", "https://outra.com.br") == ""         # não achou


def test_pasta_temporaria_que_sobrou_nao_impede_a_coleta(tmp_path, monkeypatch):
    """13/10: '_coletando_1001' ficou no OneDrive e a coleta seguinte deu FileExistsError."""
    temporaria = tmp_path / "_coletando_1001"
    temporaria.mkdir()
    (temporaria / "1001_GO_1.png").write_bytes(b"sobra")
    monkeypatch.setattr(coletor.shutil, "rmtree", lambda *a, **k: None)   # como no OneDrive
    assert coletor._preparar_temporaria(temporaria) == temporaria
    assert temporaria.is_dir() and not any(temporaria.iterdir())


def test_resumo_com_cabecalho_antigo_e_insumo_de_outro_grupo(tmp_path, monkeypatch):
    """13/10: o resumo do EPI tinha 'Link Aproveitável' e os insumos 1000 e 1056."""
    import importlib.util
    import sys
    from pathlib import Path

    caminho = Path(__file__).resolve().parents[1] / "scripts" / "corrigir_entrega_2026_10_12.py"
    spec = importlib.util.spec_from_file_location("corrigir_entrega", caminho)
    script = importlib.util.module_from_spec(spec)
    monkeypatch.setitem(sys.modules, "corrigir_entrega", script)   # exigido pelos dataclasses
    spec.loader.exec_module(script)
    pasta = tmp_path / "EPI UNIFORME"
    pasta.mkdir()
    (pasta / "resumo_coleta.csv").write_text(
        "ID FGV;Grupo;UF;Preços coletados;Faltam;Link Aproveitável\n"
        "1000;LIMPEZA E HIGIENE;GO;0;5;0\n1104;EPI / UNIFORME;GO;3;2;2\n", encoding="utf-8-sig")
    monkeypatch.setattr(script, "DIR_ENTREGA", tmp_path)
    linha = [None] * 16
    linha[1] = "1104"
    script._gravar_resumo("EPI UNIFORME", [linha, linha.copy()])
    texto = (pasta / "resumo_coleta.csv").read_text(encoding="utf-8-sig").splitlines()
    assert texto == ["ID FGV;Grupo;UF;Preços coletados;Faltam;Sem preço visível (coluna P)",
                     "1104;EPI / UNIFORME;GO;2;3;0"]
