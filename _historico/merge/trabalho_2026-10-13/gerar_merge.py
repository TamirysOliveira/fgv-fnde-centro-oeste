"""Gera fnde_resultado_merge.xlsx por grupo (13/10/2026): linhas da Entrega + linhas do William
aprovadas na revisão dos prints (decisoes_prints.json). Não altera FNDE_output.xlsx."""
import json, re, io, copy, csv, shutil, datetime, glob
from pathlib import Path
import openpyxl, pandas as pd
from PIL import Image
BASE=Path("FNDE/Centro Oeste"); ENT=BASE/"Entrega"; HIST=Path("_historico/merge_william_2026-10-13")
dig=lambda c: re.sub(r"\D","",str(c))
dec=[d for d in json.load(open(HIST/"decisoes_prints.json",encoding="utf-8")) if d["status"]=="SELECIONADA"]
receita=json.load(open(HIST.parent/"receita_fornecedores_william_2026-10-13.json",encoding="utf-8"))
# UF: planilhas de Fornecedores (coluna Estado), senão Receita
uf={}
for f in glob.glob(str(BASE/"Fornecedores/FNDE_*.xlsx")):
    for sh,v in pd.read_excel(f,sheet_name=None,dtype=str).items():
        if "CNPJ" in v:
            for _,r in v.iterrows():
                if pd.notna(r["CNPJ"]) and pd.notna(r["Estado"]): uf.setdefault(dig(r["CNPJ"]),set()).add(r["Estado"].strip())
def uf_de(c):
    s=uf.get(dig(c))
    if s:
        assert len(s)==1,(c,s); return next(iter(s))
    return receita[dig(c)][1]
ARQ={"EPIS":"[william] EPIS_E_UNIFORMES 1.xlsx","MATE":"[william] MATERIAIS_DE_CONSUMO 1.xlsx","FERR":"[william] FERRAMENTAS_E_EQUIPAMENTOS 1.xlsx"}
GRUPO={"EPIS":("EPI UNIFORME","EPI / UNIFORME","EPI_UNIFORME"),"MATE":("LIMPEZA E HIGIENE","LIMPEZA E HIGIENE","LIMPEZA E HIGIENE"),"FERR":("LIMPEZA PREDIAL","LIMPEZA PREDIAL","LIMPEZA PREDIAL")}
wbs={k:openpyxl.load_workbook(BASE/v) for k,v in ARQ.items()}
imgs={k:{im.anchor._from.row+1:im for im in wbs[k].active._images} for k in wbs}
log=[]
INS=pd.read_csv(BASE/"Pesquisa_SEAD_GO_Limpeza_Predial.csv",sep=";",dtype=str,encoding="cp1252").set_index("ID FGV")
_m=pd.read_excel(ENT/"LIMPEZA E HIGIENE/FNDE_output.xlsx",dtype=str)
assert all(INS.loc[i,"Descrição"]==t for i,t in zip(_m["Código FGV"],_m["Item"])), "Item difere da Descrição"
for pref,(pasta,cat,sub) in GRUPO.items():
    src=ENT/pasta/"FNDE_output.xlsx"; wb=openpyxl.load_workbook(src); ws=wb.active
    head=[c.value for c in ws[1]]
    meus=[[c.value for c in row] for row in ws.iter_rows(min_row=2)]
    item={str(r[1]):r[2] for r in meus}
    estilo=[copy.copy(c._style) for c in ws[2]]
    wsw=wbs[pref].active
    novos=[]
    for d in [d for d in dec if d["pref"]==pref]:
        v=[c.value for c in wsw[d["linha"]]]
        idf=str(v[1]); assert idf==d["id"]
        go = uf_de(v[6])=="GO"
        frete=None if go else v[13]
        assert go or (frete is not None and frete>0), d
        obs="Frete grátis (GO para GO)" if go else f"Opção mais barata para o CEP 74933-130: {v[14]}"
        data=datetime.datetime.strptime(str(v[4])[:10],"%Y-%m-%d")
        if idf not in item: item[idf]=INS.loc[idf,"Descrição"]
        p=d["preco_regra"]
        novos.append(dict(id=idf,row=[cat,idf,item[idf],uf_de(v[6]),data,v[5],v[6],v[7],v[8],p,None,None,p,frete,obs,None],img=imgs[pref][d["linha"]],w=d))
    # junta: por ID, primeiro as suas linhas e depois as do William
    ordem=sorted({str(r[1]) for r in meus}|{n["id"] for n in novos},key=int)
    linhas=[]
    for i in ordem:
        linhas += [("meu",r) for r in meus if str(r[1])==i]
        linhas += [("w",n) for n in novos if n["id"]==i]
    # prints: continua a numeração existente
    pdir=ENT/pasta/"prints"/sub; pdir.mkdir(parents=True,exist_ok=True)
    out=wb; wo=ws
    wo.auto_filter.ref=None
    wo.delete_rows(2, wo.max_row)
    cont={}
    for orig,x in linhas:
        vals=x if orig=="meu" else x["row"]
        wo.append(vals); r=wo.max_row
        for c,s in zip(wo[r],estilo): c._style=copy.copy(s)
        h=wo.cell(r,8)
        if h.value: h.hyperlink=h.value
        if orig=="w":
            idf=x["id"]
            if idf not in cont: cont[idf]=sum(1 for r in meus if str(r[1])==idf and not r[15])
            cont[idf]+=1; n=cont[idf]; dest=pdir/f"{idf}_GO_{n}.png"
            assert not any(str(r[1])==idf for r in []), ""
            Image.open(io.BytesIO(x["img"]._data())).save(dest)
            log.append([pasta,idf,x["row"][5],x["row"][6],x["row"][3],x["row"][9],x["row"][13],str(dest),f'{ARQ[pref]} linha {x["w"]["linha"]}'])
    wo.auto_filter.ref=f"A1:P{wo.max_row}"
    dst=ENT/pasta/"fnde_resultado_merge.xlsx"; out.save(dst)
    print(pasta,"suas",len(meus),"william",len(novos),"total",wo.max_row-1)
with open(HIST/"linhas_william_incluidas.csv","w",encoding="utf-8-sig",newline="") as fh:
    w=csv.writer(fh,delimiter=";"); w.writerow(["Grupo","ID FGV","Empresa","CNPJ","UF","Preço","Frete","Print gerado","Origem"]); w.writerows(log)
