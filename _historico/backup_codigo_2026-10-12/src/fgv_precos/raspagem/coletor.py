"""
Orquestra a coleta: para cada insumo, percorre os fornecedores do grupo até
juntar 5 preços com frete para GO.

Regras (ver ``config_raspagem.py``):
    * Ordem dos fornecedores = ordem da aba "Fornecedores" do
      FNDE_<Grupo>.xlsx (os de GO/DF vêm primeiro).
    * Produto correspondente sem preço visível (login, sob consulta ou motivo
      não identificado) -> linha para coleta manual com o link na coluna P e
      o motivo em OBS FRETE; não conta entre os 5 preços e o próximo
      fornecedor é tentado. Só o produto claramente esgotado é ignorado.
    * Produto aceito: melhor pontuação de ``descricao.pontuar`` (> 0),
      disponível e com preço.
    * Frete que exige login/cadastro -> linha na coluna P (link para coleta
      manual), com "Frete só com login" em OBS FRETE; não conta entre os 5.
    * Frete que não pode ser calculado por outro motivo -> fornecedor
      descartado.
"""

from __future__ import annotations

import logging
from dataclasses import replace
from datetime import date
from pathlib import Path

import openpyxl
import pandas as pd
from playwright.sync_api import Error as PlaywrightError
from playwright.sync_api import sync_playwright

from fgv_precos.fornecedores.insumos import ler_insumos
from fgv_precos.raspagem.config_raspagem import (
    ARQ_INSUMOS,
    ARQUIVO_FORNECEDORES_POR_GRUPO,
    CEPS,
    DIR_PLANILHAS_FORNECEDORES,
    DIR_PRINTS,
    PASTA_PRINT_POR_GRUPO,
    PRECOS_POR_UF,
    TIMEOUT_NAVEGACAO_MS,
    VIEWPORT,
)
from fgv_precos.raspagem.descricao import Insumo, interpretar, pontuar
from fgv_precos.raspagem.lojas import loja_para
from fgv_precos.raspagem.modelos import Coleta, Fornecedor, float_para_br
from fgv_precos.raspagem.plataformas.base import MOTIVO_FRETE_LOGIN

log = logging.getLogger(__name__)


def ler_fornecedores(grupo: str) -> list[Fornecedor]:
    """Lê a aba 'Fornecedores' do FNDE_<Grupo>.xlsx."""
    arquivo = DIR_PLANILHAS_FORNECEDORES / ARQUIVO_FORNECEDORES_POR_GRUPO[grupo]
    if not arquivo.exists():
        log.error("Planilha de fornecedores não encontrada: %s", arquivo)
        return []
    wb = openpyxl.load_workbook(arquivo, read_only=True, data_only=True)
    if "Fornecedores" not in wb.sheetnames:
        log.error("Aba 'Fornecedores' ausente em %s", arquivo.name)
        return []
    linhas = list(wb["Fornecedores"].iter_rows(values_only=True))
    if not linhas:
        return []
    cab = [str(c or "").strip() for c in linhas[0]]
    idx = {nome: cab.index(nome) for nome in ("Nome da Empresa", "CNPJ", "Site Oficial", "Estado")
           if nome in cab}
    saida = []
    for linha in linhas[1:]:
        site = str(linha[idx["Site Oficial"]] or "").strip() if "Site Oficial" in idx else ""
        nome = str(linha[idx["Nome da Empresa"]] or "").strip()
        if not nome or not site:
            if nome:
                log.info("Fornecedor sem site ignorado: %s", nome)
            continue
        saida.append(Fornecedor(
            nome=nome, cnpj=str(linha[idx["CNPJ"]] or "").strip() if "CNPJ" in idx else "",
            site=site, estado=str(linha[idx["Estado"]] or "").strip() if "Estado" in idx else ""))
    return saida


def ler_insumos_selecionados(ids: list[str] | None, caminho: Path = ARQ_INSUMOS) -> list[Insumo]:
    df = ler_insumos(caminho)
    if ids:
        df = df[df["ID FGV"].astype(str).isin([str(i) for i in ids])]
        faltando = set(map(str, ids)) - set(df["ID FGV"].astype(str))
        if faltando:
            log.warning("IDs não encontrados na planilha de insumos: %s", sorted(faltando))
    return [interpretar(r["ID FGV"], r["Grupo de Insumo"], r["Descrição"], r.get("Embalagem", ""))
            for _, r in df.iterrows()]


# Alguns sites recusam conexões do Chromium "de teste" do Playwright
# (erro "conexão interrompida" no log de 09/10/2026). Por isso o padrão é usar
# o Edge ou o Chrome instalados no computador, que usam a configuração de rede
# do Windows, e só cair no Chromium do Playwright se nenhum deles existir.
CANAIS_NAVEGADOR = ("msedge", "chrome", None)
ARGS_NAVEGADOR = ["--disable-blink-features=AutomationControlled"]
USER_AGENT = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
              "(KHTML, like Gecko) Chrome/130.0.0.0 Safari/537.36 Edg/130.0.0.0")


def abrir_navegador(pw, headless: bool):
    """Abre o primeiro navegador disponível da lista ``CANAIS_NAVEGADOR``."""
    ultimo_erro = None
    for canal in CANAIS_NAVEGADOR:
        try:
            navegador = pw.chromium.launch(channel=canal, headless=headless, args=ARGS_NAVEGADOR) \
                if canal else pw.chromium.launch(headless=headless, args=ARGS_NAVEGADOR)
            log.info("Navegador usado: %s", canal or "chromium do Playwright")
            return navegador
        except PlaywrightError as exc:
            ultimo_erro = exc
            log.info("Navegador '%s' indisponível: %s", canal, str(exc).splitlines()[0])
    raise RuntimeError(f"Nenhum navegador pôde ser aberto: {ultimo_erro}")


class Coletor:
    """Mantém um navegador aberto e coleta insumo a insumo."""

    def __init__(self, headless: bool = True) -> None:
        self.headless = headless
        self.data_coleta = date.today().strftime("%d/%m/%Y")

    def coletar(self, insumos: list[Insumo]) -> tuple[list[Coleta], list[dict]]:
        coletas: list[Coleta] = []
        resumo: list[dict] = []
        with sync_playwright() as pw:
            navegador = abrir_navegador(pw, self.headless)
            contexto = navegador.new_context(viewport=VIEWPORT, locale="pt-BR",
                                             timezone_id="America/Sao_Paulo",
                                             user_agent=USER_AGENT)
            contexto.set_default_timeout(TIMEOUT_NAVEGACAO_MS)
            page = contexto.new_page()
            try:
                for insumo in insumos:
                    if insumo.grupo not in ARQUIVO_FORNECEDORES_POR_GRUPO:
                        log.warning("Grupo sem planilha de fornecedores: %s", insumo.grupo)
                        continue
                    novas = self._coletar_insumo(page, insumo)
                    coletas.extend(novas)
                    for uf in CEPS:
                        n = sum(1 for c in novas if c.uf == uf and not c.link_aproveitavel)
                        manual = sum(1 for c in novas if c.uf == uf and c.link_aproveitavel)
                        resumo.append({"ID FGV": insumo.id_fgv, "Grupo": insumo.grupo, "UF": uf,
                                       "Preços coletados": n,
                                       "Faltam": max(PRECOS_POR_UF - n, 0),
                                       "Sem preço visível (coluna P)": manual})
            finally:
                contexto.close()
                navegador.close()
        return coletas, resumo

    def _coletar_insumo(self, page, insumo: Insumo) -> list[Coleta]:
        log.info("=== Insumo %s (%s): busca %s | medidas %s",
                 insumo.id_fgv, insumo.grupo, insumo.nomes_busca,
                 sorted(insumo.medidas_obrigatorias) or "-")
        pasta = DIR_PRINTS / PASTA_PRINT_POR_GRUPO[insumo.grupo]
        pasta.mkdir(parents=True, exist_ok=True)
        # Nova coleta do insumo substitui a anterior: apaga os prints antigos dele
        # (só os deste ID; prints de outros insumos ficam intactos).
        for antigo in pasta.glob(f"{insumo.id_fgv}_*.png"):
            try:
                antigo.unlink()
            except OSError as exc:
                log.warning("Não foi possível apagar o print antigo %s: %s", antigo.name, exc)
        contagem = {uf: 0 for uf in CEPS}
        coletas: list[Coleta] = []

        for fornecedor in ler_fornecedores(insumo.grupo):
            if all(n >= PRECOS_POR_UF for n in contagem.values()):
                break
            # "zeusdobrasil.com.br / epizeus.com.br": tenta cada site até achar o produto.
            for site in [s.strip() for s in fornecedor.site.split(" / ") if s.strip()]:
                f_site = replace(fornecedor, site=site)
                try:
                    novas = self._tentar_fornecedor(page, insumo, f_site, contagem, pasta)
                except PlaywrightError as exc:
                    log.error("[%s] Erro de navegador: %s", fornecedor.nome, exc)
                    novas = []
                except Exception:  # noqa: BLE001 - um site com problema não pode parar a coleta
                    log.exception("[%s] Erro inesperado", fornecedor.nome)
                    novas = []
                coletas.extend(novas)
                if any(not c.link_aproveitavel for c in novas):
                    break           # achou preço neste site; não precisa do próximo

        for uf, n in contagem.items():
            nivel = logging.INFO if n >= PRECOS_POR_UF else logging.WARNING
            log.log(nivel, "Insumo %s - %s: %d de %d preços.", insumo.id_fgv, uf, n, PRECOS_POR_UF)
        return coletas

    def _tentar_fornecedor(self, page, insumo, fornecedor, contagem, pasta) -> list[Coleta]:
        loja = loja_para(page, fornecedor)
        loja.insumo = insumo
        log.info("[%s] %s (%s)", fornecedor.nome, loja.base, loja.plataforma)

        # 1) busca: tenta cada alternativa do nome (e, por fim, só a palavra
        #    principal) até achar candidatos aderentes
        candidatos = []
        for termo in insumo.termos_de_busca:
            achados = loja.buscar(termo)
            for c in achados:
                c.pontuacao = pontuar(insumo, c.nome)
            candidatos = sorted((c for c in achados if c.pontuacao > 0),
                                key=lambda c: c.pontuacao, reverse=True)
            log.info("[%s] Busca '%s': %d resultados, %d aderentes",
                     fornecedor.nome, termo, len(achados), len(candidatos))
            if candidatos:
                break
        if not candidatos:
            return []

        # 2) primeiro candidato disponível e com preço
        produto = None
        for cand in candidatos[:5]:
            p = loja.ler_produto(cand.url)
            if p and p.disponivel and p.motivo_sem_preco and pontuar(insumo, p.nome or cand.nome) > 0:
                # Sem preço visível: registra o link para coleta manual (coluna P).
                log.info("[%s] %s: %s", fornecedor.nome, p.motivo_sem_preco, p.url)
                uf = next(iter(CEPS))
                return [Coleta(insumo.id_fgv, insumo.grupo, insumo.descricao, uf,
                               self.data_coleta, fornecedor, p, None, link_aproveitavel=p.url)]
            if p and p.disponivel and p.preco_final and pontuar(insumo, p.nome or cand.nome) > 0:
                produto = p
                break
            log.info("[%s] Descartado '%s' (esgotado ou não aderente)",
                     fornecedor.nome, cand.nome)
        if produto is None:
            return []
        log.info("[%s] Produto: %s | R$ %.2f", fornecedor.nome, produto.nome, produto.preco_final)

        # 3) frete + print para cada UF que ainda precisa de preço
        coletas = []
        for uf, cep in CEPS.items():
            if contagem[uf] >= PRECOS_POR_UF:
                continue
            if page.url != produto.url and not loja._ir(produto.url):
                continue
            frete = loja.calcular_frete(cep, uf)
            if not frete.ok and frete.motivo == MOTIVO_FRETE_LOGIN:
                # Regra (09/10/2026): frete/carrinho que exige login vai para a
                # coluna P ("Link aproveitável"), para coleta manual; não conta
                # entre os 5 preços e o próximo fornecedor é tentado.
                motivo = f"Frete só com login (preço na página: R$ {float_para_br(produto.preco_final)})"
                log.info("[%s] %s: %s", fornecedor.nome, motivo, produto.url)
                coletas.append(Coleta(insumo.id_fgv, insumo.grupo, insumo.descricao, uf,
                                      self.data_coleta, fornecedor,
                                      replace(produto, motivo_sem_preco=motivo), None,
                                      link_aproveitavel=produto.url))
                continue
            if not frete.ok:
                log.info("[%s] Descartado para %s: %s", fornecedor.nome, uf, frete.motivo)
                continue
            ordem = contagem[uf] + 1
            caminho = pasta / f"{insumo.id_fgv}_{uf}_{ordem}.png"
            if not loja.tirar_print(str(caminho)):
                continue
            contagem[uf] = ordem
            coletas.append(Coleta(insumo.id_fgv, insumo.grupo, insumo.descricao, uf,
                                  self.data_coleta, fornecedor, produto, frete, str(caminho)))
            log.info("[%s] %s OK: frete %s -> %s", fornecedor.nome, uf,
                     "grátis" if frete.gratis else f"R$ {frete.valor:.2f}", caminho.name)
        return coletas


def resumo_para_df(resumo: list[dict]) -> pd.DataFrame:
    return pd.DataFrame(resumo, columns=["ID FGV", "Grupo", "UF", "Preços coletados", "Faltam",
                                         "Sem preço visível (coluna P)"])
