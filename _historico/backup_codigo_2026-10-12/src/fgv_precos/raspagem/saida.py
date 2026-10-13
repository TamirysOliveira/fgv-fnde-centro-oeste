"""
Gravação do resultado em Excel (``Entrega/FNDE_output.xlsx``), no layout de
FNDE/Centro Oeste/FNDE_output.csv (o cabeçalho é lido do próprio modelo).

Cada execução sobrescreve as linhas dos insumos coletados nela e mantém as
linhas dos demais insumos já presentes no arquivo.

Mapeamento (confirmado com a equipe):
    A Categoria          = Grupo de Insumo
    B Código FGV         = ID FGV
    C Item               = Descrição completa do insumo
    D UF                 = GO
    E Data coleta        = data da execução
    F Nome Empresa       = da planilha de fornecedores
    G CNPJ               = da planilha de fornecedores
    H FONTE              = link (clicável) da página do produto de onde saiu o print
    I PRODUTO PESQUISADO = nome do produto no site
    J Preço PRODUTO      = preço visível na página (sem cupom/PIX)
    K VALOR DESCONTO     = vazio (descontos não são considerados)
    L OBS Desconto       = vazio
    M PREÇO FINAL        = igual ao Preço PRODUTO
    N VALOR FRETE        = opção de entrega mais barata; vazio se grátis
    O OBS FRETE          = "Frete grátis..." ou a modalidade escolhida + CEP;
                           nas linhas da coluna P, o motivo da falta de preço
    P Link aproveitável  = só quando o produto correspondente não mostra preço
                           (login, sob consulta ou motivo não identificado):
                           link do produto para coleta manual; colunas de
                           preço/frete ficam vazias
"""

from __future__ import annotations

import csv
import logging
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from fgv_precos.raspagem.config_raspagem import (
    ARQ_MODELO_SAIDA,
    CEPS,
    CODIFICACAO_SAIDA,
    DIR_ENTREGA,
    SEPARADOR_SAIDA,
)
from fgv_precos.raspagem.modelos import Coleta

log = logging.getLogger(__name__)

N_COLUNAS_ESPERADAS = 16
COL_FONTE, COL_LINK_APROVEITAVEL = 8, 16          # H e P (1 = A)
COLUNAS_VALOR = (10, 11, 13, 14)           # J, K, M, N: números no Excel
_FONTE_LINK = Font(name="Arial", size=10, color="0563C1", underline="single")
_FONTE = Font(name="Arial", size=10)


def colunas_modelo(caminho: Path | None = None) -> list[str]:
    """Lê o cabeçalho do modelo; assim a saída segue sempre o arquivo oficial."""
    with open(caminho or ARQ_MODELO_SAIDA, encoding=CODIFICACAO_SAIDA, newline="") as fh:
        return next(csv.reader(fh, delimiter=SEPARADOR_SAIDA))


def _linha(c: Coleta) -> list:
    base = [c.grupo, c.insumo_id, c.descricao, c.uf, c.data, c.fornecedor.nome,
            c.fornecedor.cnpj]
    p = c.produto
    if c.link_aproveitavel:
        # Sem preço visível: nada de preço/frete; motivo em OBS FRETE, link na coluna P.
        return base + ["", p.nome, None, None, None, None, None, p.motivo_sem_preco,
                       c.link_aproveitavel]
    f = c.frete
    cep = f"{CEPS[c.uf][:5]}-{CEPS[c.uf][5:]}"
    if f.gratis and f.modalidade.startswith("Frete grátis informado"):
        obs_frete = f"{f.modalidade} (sem calculadora de CEP no site)"
    elif f.gratis:
        obs_frete = f"Frete grátis para o CEP {cep}"
    else:
        obs_frete = f"Opção mais barata para o CEP {cep}: {f.modalidade}"
    preco = round(p.preco_final, 2)
    return base + [p.url, p.nome, preco, None, None, preco,
                   None if f.gratis else round(f.valor, 2), obs_frete, None]


ARQ_SAIDA_NOME = "FNDE_output.xlsx"


def _ler_linhas_existentes(arquivo: Path, n_colunas: int) -> list[list]:
    """Linhas (sem cabeçalho) de uma saída anterior, para preservar outros insumos."""
    if not arquivo.exists():
        return []
    try:
        ws = load_workbook(arquivo, read_only=True).active
        return [list(r[:n_colunas]) for r in ws.iter_rows(min_row=2, values_only=True)
                if any(v not in (None, "") for v in r)]
    except Exception as exc:  # noqa: BLE001 - arquivo corrompido não pode parar a gravação
        log.warning("Não foi possível ler a saída anterior %s: %s", arquivo.name, exc)
        return []


def _escrever_linha(ws, valores: list, n_colunas: int) -> None:
    ws.append(valores)
    r = ws.max_row
    for col in range(1, n_colunas + 1):
        cel = ws.cell(row=r, column=col)
        cel.font = _FONTE
        cel.alignment = Alignment(vertical="top", wrap_text=col in (3, 9, 15))
        if col in COLUNAS_VALOR and isinstance(cel.value, (int, float)):
            cel.number_format = "#,##0.00"
    for col in (COL_FONTE, COL_LINK_APROVEITAVEL):
        cel = ws.cell(row=r, column=col)
        if str(cel.value or "").startswith("http"):
            cel.hyperlink = cel.value
            cel.font = _FONTE_LINK


def gravar(coletas: list[Coleta], ids_coletados: list[str], destino: Path = DIR_ENTREGA) -> Path:
    """Grava/atualiza ``Entrega/FNDE_output.xlsx``.

    As linhas dos insumos coletados nesta execução SUBSTITUEM as linhas desses
    mesmos insumos na saída anterior; as linhas dos demais insumos são mantidas.
    """
    cab = colunas_modelo()
    if len(cab) != N_COLUNAS_ESPERADAS:
        log.warning("Modelo FNDE_output.csv tem %d colunas (esperado %d): %s",
                    len(cab), N_COLUNAS_ESPERADAS, cab)
    destino.mkdir(parents=True, exist_ok=True)
    arquivo = destino / ARQ_SAIDA_NOME

    ids = {str(i) for i in ids_coletados}
    mantidas = [l for l in _ler_linhas_existentes(arquivo, len(cab)) if str(l[1]) not in ids]
    novas = [_linha(c) for c in sorted(coletas, key=lambda c: (c.insumo_id, c.uf,
                                                                bool(c.link_aproveitavel)))]
    # Ordena por código; dentro do insumo, preços completos antes das linhas da coluna P.
    todas = sorted(mantidas + novas, key=lambda l: (str(l[1]), bool(l[COL_LINK_APROVEITAVEL - 1])))

    wb = Workbook()
    ws = wb.active
    ws.title = "FNDE_output"
    ws.append(cab)
    for cel in ws[1]:
        cel.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cel.fill = PatternFill("solid", start_color="1F4E78")
        cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for linha in todas:
        _escrever_linha(ws, linha, len(cab))

    larguras = [18, 11, 45, 5, 11, 32, 20, 40, 40, 13, 12, 12, 12, 12, 45, 40]
    for i, w in enumerate(larguras[: len(cab)], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cab))}{ws.max_row}"
    try:
        wb.save(arquivo)
    except PermissionError as exc:
        raise PermissionError(f"Feche o arquivo no Excel e rode de novo: {arquivo}") from exc
    log.info("Saída gravada: %s (%d linhas novas, %d de outros insumos mantidas)",
             arquivo, len(novas), len(mantidas))
    return arquivo
